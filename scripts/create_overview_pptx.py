"""Generate a 2-slide overview PowerPoint for the Visio Azure MCP project."""

from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pathlib import Path

# ── Colours ───────────────────────────────────────────────────────
AZURE_BLUE = RGBColor(0x00, 0x78, 0xD4)
DARK_BLUE = RGBColor(0x00, 0x2B, 0x49)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
LIGHT_GRAY = RGBColor(0xF2, 0xF2, 0xF2)
ACCENT_GREEN = RGBColor(0x10, 0x7C, 0x10)
ACCENT_ORANGE = RGBColor(0xFF, 0x8C, 0x00)
SUBTLE_GRAY = RGBColor(0x60, 0x60, 0x60)
BORDER_GRAY = RGBColor(0xD0, 0xD0, 0xD0)

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)


def _add_bg(slide, color):
    bg = slide.background
    fill = bg.fill
    fill.solid()
    fill.fore_color.rgb = color


def _add_rect(slide, left, top, width, height, fill_color, line_color=None):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill_color
    if line_color:
        shape.line.color.rgb = line_color
        shape.line.width = Pt(1)
    else:
        shape.line.fill.background()
    shape.shadow.inherit = False
    return shape


def _add_text_box(slide, left, top, width, height, text, font_size=12,
                  bold=False, color=DARK_BLUE, alignment=PP_ALIGN.LEFT, font_name="Segoe UI"):
    txBox = slide.shapes.add_textbox(left, top, width, height)
    tf = txBox.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = text
    p.font.size = Pt(font_size)
    p.font.bold = bold
    p.font.color.rgb = color
    p.font.name = font_name
    p.alignment = alignment
    return txBox


def _add_bullet_list(slide, left, top, width, height, items, font_size=11,
                     color=DARK_BLUE, bold_prefix=True):
    txBox = slide.shapes.add_textbox(left, top, width, height)
    tf = txBox.text_frame
    tf.word_wrap = True
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_before = Pt(3)
        p.space_after = Pt(1)

        if bold_prefix and " – " in item:
            bold_part, rest = item.split(" – ", 1)
            run_b = p.add_run()
            run_b.text = "• " + bold_part + " – "
            run_b.font.size = Pt(font_size)
            run_b.font.bold = True
            run_b.font.color.rgb = color
            run_b.font.name = "Segoe UI"
            run_r = p.add_run()
            run_r.text = rest
            run_r.font.size = Pt(font_size)
            run_r.font.bold = False
            run_r.font.color.rgb = color
            run_r.font.name = "Segoe UI"
        else:
            run = p.add_run()
            run.text = "• " + item
            run.font.size = Pt(font_size)
            run.font.color.rgb = color
            run.font.name = "Segoe UI"
    return txBox


def _add_icon_label(slide, left, top, icon_text, label, desc, icon_color=AZURE_BLUE):
    """Add an icon circle with label and description."""
    # Icon circle
    circle = slide.shapes.add_shape(MSO_SHAPE.OVAL, left, top, Inches(0.5), Inches(0.5))
    circle.fill.solid()
    circle.fill.fore_color.rgb = icon_color
    circle.line.fill.background()
    circle.shadow.inherit = False
    tf = circle.text_frame
    tf.word_wrap = False
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    run = p.add_run()
    run.text = icon_text
    run.font.size = Pt(14)
    run.font.color.rgb = WHITE
    run.font.bold = True
    run.font.name = "Segoe UI"
    tf.paragraphs[0].space_before = Pt(0)
    tf.paragraphs[0].space_after = Pt(0)

    # Label
    _add_text_box(slide, left + Inches(0.6), top - Inches(0.02), Inches(2.0), Inches(0.3),
                  label, font_size=12, bold=True, color=DARK_BLUE)
    # Desc
    _add_text_box(slide, left + Inches(0.6), top + Inches(0.22), Inches(2.2), Inches(0.5),
                  desc, font_size=9, color=SUBTLE_GRAY)


# ══════════════════════════════════════════════════════════════════
# SLIDE 1 — Overview & Architecture
# ══════════════════════════════════════════════════════════════════
slide1 = prs.slides.add_slide(prs.slide_layouts[6])  # blank
_add_bg(slide1, WHITE)

# Title bar
title_bar = _add_rect(slide1, Inches(0), Inches(0), Inches(13.333), Inches(1.1), AZURE_BLUE)
_add_text_box(slide1, Inches(0.6), Inches(0.15), Inches(9), Inches(0.5),
              "Visio Azure MCP", font_size=32, bold=True, color=WHITE)
_add_text_box(slide1, Inches(0.6), Inches(0.6), Inches(10), Inches(0.4),
              "AI-Powered Azure Architecture Diagrams  •  MCP Server  •  Streamlit App  •  VS Code Extension  •  Desktop App",
              font_size=13, color=RGBColor(0xCC, 0xE4, 0xFF))

# Stats bar
_add_rect(slide1, Inches(9.5), Inches(0.15), Inches(3.5), Inches(0.8), RGBColor(0x00, 0x5A, 0x9E))
stats = [("28", "Tools"), ("8", "Resources"), ("7", "Prompts"), ("206", "Catalog")]
for i, (num, label) in enumerate(stats):
    x = Inches(9.65) + Inches(i * 0.87)
    _add_text_box(slide1, x, Inches(0.15), Inches(0.8), Inches(0.35),
                  num, font_size=22, bold=True, color=WHITE, alignment=PP_ALIGN.CENTER)
    _add_text_box(slide1, x, Inches(0.48), Inches(0.8), Inches(0.25),
                  label, font_size=9, color=RGBColor(0xA0, 0xCC, 0xEE), alignment=PP_ALIGN.CENTER)

# ── Left column: What It Does ────────────────────────────────────
_add_text_box(slide1, Inches(0.5), Inches(1.3), Inches(4), Inches(0.35),
              "What It Does", font_size=18, bold=True, color=AZURE_BLUE)

left_items = [
    "Natural Language → Diagram – Describe an architecture in plain English; the AI builds it step by step",
    "Business Requirements → Architecture – Input a business case; AI analyses workload, selects style, picks services, builds & validates",
    "123 Azure Resource Types – Official SVG icons from Azure Public Service, Entra, and Fabric icon packs",
    "5 Reference Architecture Templates – Hand-tuned from Azure Architecture Center (Landing Zone, Web App, AKS, AI Chat, AI Landing Zone)",
    "206-Entry Architecture Catalog – Browse/search real architectures from Azure Architecture Center",
    "36 Design Patterns + 6 Styles – Cloud patterns (CQRS, Saga, Circuit Breaker…) and architecture styles (N-Tier, Microservices…)",
    "WAF + CAF Validation – Well-Architected Framework (5 pillars) and Cloud Adoption Framework (7 principles) scoring",
    "Multi-Format Output – Save as .vsdx (Visio COM) or .drawio (no Windows/Visio needed)",
    "Import & Convert – Import existing .vsdx (multi-page) or convert whiteboard photos/screenshots via AI vision",
]
_add_bullet_list(slide1, Inches(0.5), Inches(1.7), Inches(5.5), Inches(5.0),
                 left_items, font_size=10.5, color=DARK_BLUE)

# ── Right column: Architecture diagram ───────────────────────────
_add_text_box(slide1, Inches(6.6), Inches(1.3), Inches(4), Inches(0.35),
              "How It Works", font_size=18, bold=True, color=AZURE_BLUE)

# Architecture boxes
arch_y = Inches(1.8)

# User box
_add_rect(slide1, Inches(9.0), arch_y, Inches(2.0), Inches(0.55), LIGHT_GRAY, BORDER_GRAY)
_add_text_box(slide1, Inches(9.0), arch_y + Inches(0.05), Inches(2.0), Inches(0.45),
              "👤 User", font_size=12, bold=True, color=DARK_BLUE, alignment=PP_ALIGN.CENTER)

# Arrow
_add_text_box(slide1, Inches(9.8), arch_y + Inches(0.55), Inches(0.5), Inches(0.3),
              "▼", font_size=14, color=AZURE_BLUE, alignment=PP_ALIGN.CENTER)

# Streamlit App box
app_y = arch_y + Inches(0.95)
app_box = _add_rect(slide1, Inches(6.8), app_y, Inches(6.0), Inches(1.8), RGBColor(0xE8, 0xF0, 0xFE), AZURE_BLUE)
_add_text_box(slide1, Inches(7.0), app_y + Inches(0.05), Inches(5.5), Inches(0.3),
              "Streamlit App", font_size=14, bold=True, color=AZURE_BLUE)

# Inner boxes
inner_y = app_y + Inches(0.4)
# Chat
_add_rect(slide1, Inches(7.0), inner_y, Inches(1.7), Inches(1.2), WHITE, BORDER_GRAY)
_add_text_box(slide1, Inches(7.0), inner_y + Inches(0.1), Inches(1.7), Inches(0.25),
              "💬 AI Chat", font_size=10, bold=True, color=DARK_BLUE, alignment=PP_ALIGN.CENTER)
_add_text_box(slide1, Inches(7.05), inner_y + Inches(0.35), Inches(1.6), Inches(0.8),
              "Natural language\nBusiness req input\n3 AI providers", font_size=8, color=SUBTLE_GRAY, alignment=PP_ALIGN.CENTER)

# Preview
_add_rect(slide1, Inches(8.85), inner_y, Inches(1.7), Inches(1.2), WHITE, BORDER_GRAY)
_add_text_box(slide1, Inches(8.85), inner_y + Inches(0.1), Inches(1.7), Inches(0.25),
              "📊 Preview", font_size=10, bold=True, color=DARK_BLUE, alignment=PP_ALIGN.CENTER)
_add_text_box(slide1, Inches(8.9), inner_y + Inches(0.35), Inches(1.6), Inches(0.8),
              "Live SVG render\nPage tabs\nReal-time updates", font_size=8, color=SUBTLE_GRAY, alignment=PP_ALIGN.CENTER)

# Sidebar
_add_rect(slide1, Inches(10.7), inner_y, Inches(1.9), Inches(1.2), WHITE, BORDER_GRAY)
_add_text_box(slide1, Inches(10.7), inner_y + Inches(0.1), Inches(1.9), Inches(0.25),
              "📋 Sidebar", font_size=10, bold=True, color=DARK_BLUE, alignment=PP_ALIGN.CENTER)
_add_text_box(slide1, Inches(10.75), inner_y + Inches(0.35), Inches(1.8), Inches(0.8),
              "Biz → Arch input\nRef arch templates\nImport / Save\nCatalog browser", font_size=8, color=SUBTLE_GRAY, alignment=PP_ALIGN.CENTER)

# Arrow
_add_text_box(slide1, Inches(9.8), app_y + Inches(1.8), Inches(0.5), Inches(0.3),
              "▼ stdio JSON-RPC", font_size=8, color=AZURE_BLUE, alignment=PP_ALIGN.CENTER)

# MCP Server box
mcp_y = app_y + Inches(2.2)
mcp_box = _add_rect(slide1, Inches(6.8), mcp_y, Inches(6.0), Inches(2.1), RGBColor(0xE2, 0xF0, 0xD9), ACCENT_GREEN)
_add_text_box(slide1, Inches(7.0), mcp_y + Inches(0.05), Inches(5.5), Inches(0.3),
              "MCP Server (FastMCP)", font_size=14, bold=True, color=ACCENT_GREEN)

# MCP inner grid
mgrid_y = mcp_y + Inches(0.45)
mcp_modules = [
    ("Diagram State", "CRUD · layout\nassign · query"),
    ("Visio Engine", "COM automation\nSVG import"),
    ("Draw.io Engine", "mxGraph XML\n97+ icon styles"),
    ("WAF Validator", "5 pillars · smart\nmulti-region detect"),
    ("CAF Validator", "7 principles\nnaming · tagging"),
    ("Ref Architectures", "5 templates\n206 catalog"),
]
for i, (title, desc) in enumerate(mcp_modules):
    col = i % 3
    row = i // 3
    bx = Inches(7.0) + Inches(col * 1.95)
    by = mgrid_y + Inches(row * 0.75)
    _add_rect(slide1, bx, by, Inches(1.8), Inches(0.65), WHITE, BORDER_GRAY)
    _add_text_box(slide1, bx, by + Inches(0.03), Inches(1.8), Inches(0.2),
                  title, font_size=9, bold=True, color=DARK_BLUE, alignment=PP_ALIGN.CENTER)
    _add_text_box(slide1, bx, by + Inches(0.25), Inches(1.8), Inches(0.4),
                  desc, font_size=7.5, color=SUBTLE_GRAY, alignment=PP_ALIGN.CENTER)

# Footer
_add_rect(slide1, Inches(0), Inches(7.0), Inches(13.333), Inches(0.5), DARK_BLUE)
_add_text_box(slide1, Inches(0.5), Inches(7.05), Inches(12), Inches(0.35),
              "github.com/CloudDaddyZA/VisioIntegration  •  Python 3.14  •  MCP Protocol  •  Visio COM + Draw.io  •  GitHub Copilot / OpenAI / Azure OpenAI",
              font_size=10, color=RGBColor(0xA0, 0xCC, 0xEE), alignment=PP_ALIGN.CENTER)


# ══════════════════════════════════════════════════════════════════
# SLIDE 2 — Features Deep Dive
# ══════════════════════════════════════════════════════════════════
slide2 = prs.slides.add_slide(prs.slide_layouts[6])  # blank
_add_bg(slide2, WHITE)

# Title bar
_add_rect(slide2, Inches(0), Inches(0), Inches(13.333), Inches(0.9), AZURE_BLUE)
_add_text_box(slide2, Inches(0.6), Inches(0.2), Inches(9), Inches(0.45),
              "Features & Capabilities", font_size=28, bold=True, color=WHITE)

# ── Column 1: MCP Server Tools ───────────────────────────────────
col1_x = Inches(0.4)
_add_text_box(slide2, col1_x, Inches(1.1), Inches(3.8), Inches(0.3),
              "🔧  MCP Server — 28 Tools", font_size=15, bold=True, color=AZURE_BLUE)

tools_items = [
    "Diagram CRUD – create, add resource, boundary, connect, remove",
    "Auto-Layout – tiered/grid/grouped with boundary awareness",
    "Reference Archs – 5 built-in + apply/merge with existing diagrams",
    "Architecture Catalog – browse/search 206 entries from Azure Arch Center",
    "Design Knowledge – 36 patterns + 6 architecture styles with guidance",
    "Validation – WAF (5 pillars) + CAF (7 principles) scoring & findings",
    "Import – .vsdx multi-page parsing, image/screenshot AI conversion",
    "Rendering – .vsdx (Visio COM) or .drawio (mxGraph XML, no Windows needed)",
    "Shape Catalog – 123 types, 97 SVG icons, 40+ aliases, Entra + Fabric icons",
]
_add_bullet_list(slide2, col1_x, Inches(1.45), Inches(4.0), Inches(4.5),
                 tools_items, font_size=10, color=DARK_BLUE)

# ── Column 2: AI & Business Features ─────────────────────────────
col2_x = Inches(4.7)
_add_text_box(slide2, col2_x, Inches(1.1), Inches(4.0), Inches(0.3),
              "🤖  AI-Powered Workflows", font_size=15, bold=True, color=AZURE_BLUE)

ai_items = [
    "Business → Architecture – describe a business need; AI decomposes into workload analysis, architecture style selection, service picking, diagram building, WAF/CAF validation",
    "Natural Language Chat – \"Build a hub-spoke network with Azure Firewall\" → full diagram",
    "7 Prompt Templates – getting started, hub-spoke, 3-tier web, microservices, AI chat, AI landing zone, business-to-architecture",
    "14 System Guidelines – CAF naming, save workflow, merge logic, style/pattern selection, business decomposition",
    "Multi-Provider – GitHub Copilot (30+ models), OpenAI, Azure OpenAI with auto-detection",
    "Image → Diagram – upload whiteboard photo or screenshot; AI vision identifies Azure components",
]
_add_bullet_list(slide2, col2_x, Inches(1.45), Inches(4.2), Inches(4.5),
                 ai_items, font_size=10, color=DARK_BLUE)

# ── Column 3: Interfaces ─────────────────────────────────────────
col3_x = Inches(9.2)
_add_text_box(slide2, col3_x, Inches(1.1), Inches(3.8), Inches(0.3),
              "🖥️  4 Interfaces", font_size=15, bold=True, color=AZURE_BLUE)

iface_items = [
    "Streamlit Web App – chat + live SVG preview + sidebar controls on port 8501",
    "VS Code Extension – 13 commands, 3 tree views, webview preview, auto-start MCP",
    "Desktop App – PyInstaller + pywebview standalone Windows executable",
    "CLI / MCP stdio – direct MCP server for any MCP-compatible AI agent",
]
_add_bullet_list(slide2, col3_x, Inches(1.45), Inches(3.9), Inches(2.5),
                 iface_items, font_size=10, color=DARK_BLUE)

# ── Validation section ────────────────────────────────────────────
val_y = Inches(3.7)
_add_text_box(slide2, col3_x, val_y, Inches(3.8), Inches(0.3),
              "✅  Validation Engine", font_size=15, bold=True, color=ACCENT_GREEN)

_add_rect(slide2, col3_x, val_y + Inches(0.35), Inches(1.85), Inches(2.5), RGBColor(0xE8, 0xF4, 0xE8), ACCENT_GREEN)
_add_text_box(slide2, col3_x + Inches(0.1), val_y + Inches(0.4), Inches(1.65), Inches(0.25),
              "WAF (5 Pillars)", font_size=10, bold=True, color=ACCENT_GREEN)
waf_text = "• Reliability\n• Security\n• Cost Optimization\n• Operational Excellence\n• Performance Efficiency\n\nSmart detection:\n• Multi-region\n• DB failover\n• Availability zones\n• CI/CD pipelines"
_add_text_box(slide2, col3_x + Inches(0.1), val_y + Inches(0.65), Inches(1.65), Inches(2.1),
              waf_text, font_size=8, color=DARK_BLUE)

_add_rect(slide2, col3_x + Inches(1.95), val_y + Inches(0.35), Inches(1.85), Inches(2.5), RGBColor(0xE8, 0xF0, 0xFE), AZURE_BLUE)
_add_text_box(slide2, col3_x + Inches(2.05), val_y + Inches(0.4), Inches(1.65), Inches(0.25),
              "CAF (7 Principles)", font_size=10, bold=True, color=AZURE_BLUE)
caf_text = "• Naming Convention\n• Resource Organization\n• Network Topology\n• Identity & Access\n• Governance\n• Security Baseline\n• Management & Monitoring\n\n50+ resource type prefixes"
_add_text_box(slide2, col3_x + Inches(2.05), val_y + Inches(0.65), Inches(1.65), Inches(2.1),
              caf_text, font_size=8, color=DARK_BLUE)

# ── Bottom: Tech stack bar ────────────────────────────────────────
_add_rect(slide2, Inches(0.4), Inches(6.4), Inches(8.4), Inches(0.85), LIGHT_GRAY, BORDER_GRAY)
_add_text_box(slide2, Inches(0.5), Inches(6.42), Inches(8.2), Inches(0.25),
              "Tech Stack", font_size=11, bold=True, color=DARK_BLUE)
tech_labels = [
    ("Python 3.14", AZURE_BLUE),
    ("FastMCP", ACCENT_GREEN),
    ("Streamlit", RGBColor(0xFF, 0x4B, 0x4B)),
    ("OpenAI SDK", DARK_BLUE),
    ("Visio COM", RGBColor(0x7F, 0x6D, 0xB0)),
    ("Draw.io/mxGraph", ACCENT_ORANGE),
    ("PyInstaller", SUBTLE_GRAY),
    ("TypeScript/esbuild", RGBColor(0x30, 0x78, 0xC6)),
]
for i, (label, color) in enumerate(tech_labels):
    tx = Inches(0.6) + Inches(i * 1.04)
    pill = _add_rect(slide2, tx, Inches(6.72), Inches(0.95), Inches(0.35), color)
    _add_text_box(slide2, tx, Inches(6.75), Inches(0.95), Inches(0.3),
                  label, font_size=8, bold=True, color=WHITE, alignment=PP_ALIGN.CENTER)

# Footer
_add_rect(slide2, Inches(0), Inches(7.0), Inches(13.333), Inches(0.5), DARK_BLUE)
_add_text_box(slide2, Inches(0.5), Inches(7.05), Inches(12), Inches(0.35),
              "github.com/CloudDaddyZA/VisioIntegration  •  Windows 10/11  •  Visio 2019+ / M365 (optional for .drawio)  •  MIT License",
              font_size=10, color=RGBColor(0xA0, 0xCC, 0xEE), alignment=PP_ALIGN.CENTER)


# ── Save ──────────────────────────────────────────────────────────
output_dir = Path(__file__).resolve().parent.parent / "output"
output_dir.mkdir(exist_ok=True)
output_path = output_dir / "VisioAzureMCP-Overview.pptx"
prs.save(str(output_path))
print(f"Saved: {output_path}")
